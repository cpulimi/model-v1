"""Data layer for the base rebalancing model.

Everything is stored sparsely. Only strictly positive excess / need entries are
kept, so for the real ~600 sites x 60,000 parts data the dicts scale with the
number of nonzero (site, part) pairs, never with |I| x |K|.

Index conventions (match the formulation in README.md):
    i  source site     j  destination site     k  part     m  supplier

    E[(i, k)]      excess of part k at site i            (only > 0 stored)
    N[(j, k)]      need of part k at site j              (only > 0 stored)
    C[(i, j)]      ship cost per unit, same for every part        -- or --
    C[(i, j, k)]   ship cost per unit, part specific
    B[(m, j, k)]   purchase cost per unit from supplier m for (j, k)
    UD, FC1, FC2   unmet cost per unit, fixed cost per shipping site,
                   fixed cost per receiving site

Three entry points:
    generate_synthetic(...)   Table 1 of the QVA paper
    load_instance(path)       Aarav's old-model workbooks (OldModelDatasets repo,
                              <city>_<n>_hubs/<seed>.xlsx)
    write_instance_csv / read_instance_csv
                              our own round-trip format, used so every run
                              directory carries the exact instance it solved and
                              score_solution.py can re-check it without Gurobi
"""

from __future__ import annotations

import csv
import json
import warnings
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np


@dataclass
class Instance:
    sites: list
    parts: list
    suppliers: list
    E: dict  # (i, k) -> excess units, > 0 only
    N: dict  # (j, k) -> need units, > 0 only
    C: dict  # (i, j) -> cost  OR  (i, j, k) -> cost
    B: dict  # (m, j, k) -> cost
    UD: float
    FC1: float
    FC2: float
    name: str = ""
    meta: dict = field(default_factory=dict)  # generator params etc., written to meta.json

    @property
    def ship_cost_by_part(self) -> bool:
        """True when C is keyed (i, j, k), False when keyed (i, j)."""
        for key in self.C:
            return len(key) == 3
        return False

    def ship_cost(self, i, j, k):
        """Per-unit ship cost for arc (i, j, k), whichever way C is keyed. KeyError if missing."""
        if self.ship_cost_by_part:
            return self.C[(i, j, k)]
        return self.C[(i, j)]


# ---------------------------------------------------------------------------
# Sparse index sets. These define which variables exist; the solver and the
# independent scorer both use them, so they live here rather than in either.
# ---------------------------------------------------------------------------
def ship_arcs(inst: Instance) -> list:
    """All (i, j, k) with E[i,k] > 0, N[j,k] > 0, i != j.

    Built per part from the nonzero sources and sinks of that part, so the cost
    is sum_k |sources_k| * |sinks_k|, the number of arcs that really exist.
    """
    sources = defaultdict(list)
    sinks = defaultdict(list)
    for (i, k) in inst.E:
        sources[k].append(i)
    for (j, k) in inst.N:
        sinks[k].append(j)
    arcs = []
    for k, srcs in sources.items():
        dsts = sinks.get(k)
        if not dsts:
            continue
        for i in srcs:
            arcs.extend((i, j, k) for j in dsts if j != i)
    return arcs


def purchase_keys(inst: Instance) -> list:
    """All (m, j, k) with N[j,k] > 0, for every supplier m."""
    return [(m, j, k) for (j, k) in inst.N for m in inst.suppliers]


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------
def validate_instance(inst: Instance) -> list:
    """Check an instance; raise ValueError on hard errors, return (and warn) the soft ones.

    Errors:   negative E / N / C / B / UD / FC1 / FC2; non-integer E or N (the
              model has integer X, T, U and equality (3), so fractional need is
              infeasible); unknown site / part / supplier ids; a missing ship
              cost for an arc that exists; a missing purchase cost for a
              (supplier, site, part) whose need is > 0; mixed C key lengths.
    Warnings: any (site, part) with both E > 0 and N > 0. The model allows it
              (the site cannot ship to itself, i != j), but it usually means the
              upstream netting was not done.
              Any purchase cost of 0. The old Java generator sets B[m,j] = 0
              whenever m == j, so hubs 1-3 buy from suppliers 1-3 for free;
              --free-purchase block removes those T variables.
    """
    errors = []
    site_set, part_set, sup_set = set(inst.sites), set(inst.parts), set(inst.suppliers)

    for label, d in (("E", inst.E), ("N", inst.N)):
        for (s, k), v in d.items():
            if s not in site_set or k not in part_set:
                errors.append(f"{label}[{s},{k}]: unknown site or part id")
            if v < 0:
                errors.append(f"{label}[{s},{k}] = {v} is negative")
            elif v != int(v):
                errors.append(f"{label}[{s},{k}] = {v} is not an integer")

    key_lens = {len(key) for key in inst.C}
    if len(key_lens) > 1:
        errors.append(f"C mixes (i,j) and (i,j,k) keys: key lengths {sorted(key_lens)}")
    for key, v in inst.C.items():
        if v < 0:
            errors.append(f"C{list(key)} = {v} is negative")
    for key, v in inst.B.items():
        if v < 0:
            errors.append(f"B{list(key)} = {v} is negative")
        if key[0] not in sup_set:
            errors.append(f"B{list(key)}: unknown supplier id")
    for label, v in (("UD", inst.UD), ("FC1", inst.FC1), ("FC2", inst.FC2)):
        if v < 0:
            errors.append(f"{label} = {v} is negative")

    if len(key_lens) <= 1:
        missing_c = [a for a in ship_arcs(inst) if _missing_ship_cost(inst, a)]
        for a in missing_c[:20]:
            errors.append(f"no ship cost for existing arc X{list(a)}")
        if len(missing_c) > 20:
            errors.append(f"... and {len(missing_c) - 20} more arcs with no ship cost")
    missing_b = [t for t in purchase_keys(inst) if t not in inst.B]
    for t in missing_b[:20]:
        errors.append(f"no purchase cost B{list(t)} although N[{t[1]},{t[2]}] > 0")
    if len(missing_b) > 20:
        errors.append(f"... and {len(missing_b) - 20} more missing purchase costs")

    if errors:
        raise ValueError("invalid instance:\n  " + "\n  ".join(errors))

    soft = []
    both = [key for key in inst.E if key in inst.N]
    if both:
        preview = ", ".join(f"({s},{k})" for s, k in both[:10])
        soft.append(
            f"{len(both)} (site, part) pairs have both E > 0 and N > 0, e.g. {preview}. "
            "The model allows this (no self-shipping), but check the netting upstream."
        )
    free = [t for t in purchase_keys(inst) if inst.B[t] == 0]
    if free:
        preview = ", ".join(f"B[{m},{j},{k}]" for m, j, k in free[:10])
        soft.append(
            f"{len(free)} purchase costs on needed (j,k) are 0, e.g. {preview}. "
            "Those purchases are free; --free-purchase block removes them."
        )
    for msg in soft:
        warnings.warn(msg, stacklevel=2)
    return soft


def _missing_ship_cost(inst: Instance, arc) -> bool:
    try:
        inst.ship_cost(*arc)
        return False
    except KeyError:
        return True


def _drop_zeros(d: dict) -> dict:
    return {key: v for key, v in d.items() if v != 0}


# ---------------------------------------------------------------------------
# Synthetic generator (QVA paper, Table 1)
# ---------------------------------------------------------------------------
def generate_synthetic(
    n_sites: int,
    n_parts: int,
    n_suppliers: int,
    seed: int,
    fc1: float = 100.0,
    fc2: float = 1000.0,
    ud: float = 10000.0,
    ship_cost_by_part: bool = False,
) -> Instance:
    """Random instance following Table 1 of the QVA paper.

        ship cost      ~ DU(50, 150)
        purchase cost  ~ DU(75, 175)
        excess         50% of sites 0, the other 50% ~ DU(20, 120)
        need           DU(50, 150) where the site has no excess of that part, else 0
        FC1 = 100, FC2 = 1000, UD = 10000   (paper prints "102, 103, 104";
                                             read as 10^2, 10^3, 10^4)

    "50% of sites" is applied per part: for each part, exactly n_sites // 2
    sites drawn at random hold excess and the rest need it. (A per-cell coin flip
    can leave a small instance with no excess at all.)

    Ship cost is drawn per ordered site pair C[i,j] by default (same for every
    part, like a distance-based lane cost); ship_cost_by_part=True draws
    C[i,j,k] per arc instead. Purchase cost is drawn per (m, j, k) for every
    (j, k) with need. Draw order is fixed, so (sizes, seed) reproduce exactly.
    """
    if n_sites < 2:
        raise ValueError("need at least 2 sites")
    rng = np.random.default_rng(seed)
    sites = [f"s{i}" for i in range(n_sites)]
    parts = [f"p{k}" for k in range(n_parts)]
    suppliers = [f"m{m}" for m in range(n_suppliers)]
    n_excess = n_sites // 2

    E, N = {}, {}
    for k in parts:
        order = rng.permutation(n_sites)
        exc_vals = rng.integers(20, 121, size=n_excess)
        need_vals = rng.integers(50, 151, size=n_sites - n_excess)
        for idx, v in zip(order[:n_excess], exc_vals):
            E[(sites[idx], k)] = int(v)
        for idx, v in zip(order[n_excess:], need_vals):
            N[(sites[idx], k)] = int(v)

    inst = Instance(sites, parts, suppliers, E, N, C={}, B={}, UD=ud, FC1=fc1, FC2=fc2)

    if ship_cost_by_part:
        arcs = ship_arcs(inst)
        vals = rng.integers(50, 151, size=len(arcs))
        inst.C = {a: int(v) for a, v in zip(arcs, vals)}
    else:
        pairs = [(i, j) for i in sites for j in sites if i != j]
        vals = rng.integers(50, 151, size=len(pairs))
        inst.C = {p: int(v) for p, v in zip(pairs, vals)}

    tkeys = purchase_keys(inst)
    vals = rng.integers(75, 176, size=len(tkeys))
    inst.B = {t: int(v) for t, v in zip(tkeys, vals)}

    inst.name = f"syn_s{n_sites}_p{n_parts}_m{n_suppliers}_seed{seed}"
    inst.meta = {
        "source": "generate_synthetic",
        "n_sites": n_sites,
        "n_parts": n_parts,
        "n_suppliers": n_suppliers,
        "seed": seed,
        "ship_cost_by_part": ship_cost_by_part,
        "excess_rule": "per part, n_sites//2 random sites ~ DU(20,120), others 0",
        "need_rule": "DU(50,150) where E[j,k] == 0, else 0",
        "ship_cost_dist": "DU(50,150)",
        "purchase_cost_dist": "DU(75,175)",
    }
    validate_instance(inst)
    return inst


# ---------------------------------------------------------------------------
# Real data
# ---------------------------------------------------------------------------
OLD_SHEETS = {
    # sheet name -> number of index columns before the value column
    "PARA_Eik": 2,   # part k, site i, excess
    "PARA_Njk": 2,   # part k, site j, need
    "PARA_Bmjk": 3,  # part k, supplier m, site j, cost
    "PARA_Cijk": 3,  # part k, site i, site j, cost
}


def load_instance(path, ud: float = 10000.0, fc1: float = 100.0, fc2: float = 1000.0) -> Instance:
    """Load one of Aarav's old-model workbooks: <city>_<n>_hubs/<seed>.xlsx.

    Sheets (row 1 is a header, every cell is TEXT and is converted to int):
        PARA_Eik   part k, site i, excess
        PARA_Njk   part k, site j, need
        PARA_Bmjk  part k, supplier m, site j, cost
        PARA_Cijk  part k, site i, site j, cost

    Hubs play both roles (i = j); suppliers are a separate 1..3 set, so hub "1"
    and supplier "1" are different entities that happen to share a label. Ids
    are kept as the workbook's strings so names line up with Aarav's model.

    Zero E / N rows are dropped, the C diagonal (i == j) is dropped, C stays
    keyed (i, j, k). UD / FC1 / FC2 are not in the workbook and come from the
    arguments. B is kept as given, zeros included (validate_instance warns).
    """
    import openpyxl  # only needed for this loader

    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        missing = [s for s in OLD_SHEETS if s not in wb.sheetnames]
        if missing:
            raise ValueError(f"{path}: missing sheets {missing} (has {wb.sheetnames})")
        tables = {}
        for sheet, n_idx in OLD_SHEETS.items():
            rows = list(wb[sheet].iter_rows(values_only=True))
            body = {}
            for r, row in enumerate(rows[1:], start=2):
                if row is None or all(c is None or str(c).strip() == "" for c in row):
                    continue
                cells = [str(c).strip() for c in row[: n_idx + 1]]
                try:
                    value = int(cells[n_idx])
                    keys = tuple(str(int(c)) for c in cells[:n_idx])
                except (ValueError, IndexError):
                    raise ValueError(f"{path} {sheet} row {r}: not integer text: {row!r}") from None
                if keys in body:
                    raise ValueError(f"{path} {sheet} row {r}: duplicate key {keys}")
                body[keys] = value
            tables[sheet] = body
    finally:
        wb.close()

    E_raw = {(i, k): v for (k, i), v in tables["PARA_Eik"].items()}
    N_raw = {(j, k): v for (k, j), v in tables["PARA_Njk"].items()}
    B = {(m, j, k): v for (k, m, j), v in tables["PARA_Bmjk"].items()}
    C_all = {(i, j, k): v for (k, i, j), v in tables["PARA_Cijk"].items()}
    C = {key: v for key, v in C_all.items() if key[0] != key[1]}
    diag_nonzero = sum(1 for key, v in C_all.items() if key[0] == key[1] and v != 0)

    def by_number(ids):
        return sorted(ids, key=int)

    sites = by_number({s for s, _ in E_raw} | {s for s, _ in N_raw}
                      | {key[0] for key in C_all} | {key[1] for key in C_all} | {key[1] for key in B})
    parts = by_number({k for _, k in E_raw} | {k for _, k in N_raw} | {key[2] for key in B} | {key[2] for key in C_all})
    suppliers = by_number({key[0] for key in B})

    # <city>_<n>_hubs/<seed>.xlsx -> city, n, seed
    folder = path.parent.name
    city, _, rest = folder.partition("_")
    hubs_label = rest.split("_")[0] if rest else ""
    seed = path.stem

    inst = Instance(
        sites=sites,
        parts=parts,
        suppliers=suppliers,
        E=_drop_zeros(E_raw),
        N=_drop_zeros(N_raw),
        C=C,
        B=B,
        UD=ud,
        FC1=fc1,
        FC2=fc2,
        name=f"{city}_{seed}",
        meta={
            "source": "load_instance",
            "path": str(path),
            "city": city,
            "hubs_label": int(hubs_label) if hubs_label.isdigit() else None,
            "seed": int(seed) if seed.isdigit() else seed,
            "n_sites": len(sites),
            "n_parts": len(parts),
            "n_suppliers": len(suppliers),
            "rows": {s: len(t) for s, t in tables.items()},
            "zero_E_rows_dropped": sum(1 for v in E_raw.values() if v == 0),
            "zero_N_rows_dropped": sum(1 for v in N_raw.values() if v == 0),
            "C_diagonal_rows_dropped": len(C_all) - len(C),
            "C_diagonal_nonzero": diag_nonzero,
            "B_zero_entries": sum(1 for v in B.values() if v == 0),
        },
    )
    if inst.meta["hubs_label"] is not None and inst.meta["hubs_label"] != len(sites):
        warnings.warn(f"{path}: folder says {hubs_label} hubs but the workbook has {len(sites)} sites")
    validate_instance(inst)
    return inst


# ---------------------------------------------------------------------------
# Round-trip CSV format (ours, not the real-data format)
# ---------------------------------------------------------------------------
def write_instance_csv(inst: Instance, out_dir) -> Path:
    """Write inst to out_dir/{meta.json, E.csv, N.csv, C.csv, B.csv}."""
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    meta = {
        "name": inst.name,
        "sites": [str(s) for s in inst.sites],
        "parts": [str(k) for k in inst.parts],
        "suppliers": [str(m) for m in inst.suppliers],
        "UD": inst.UD,
        "FC1": inst.FC1,
        "FC2": inst.FC2,
        "ship_cost_by_part": inst.ship_cost_by_part,
        "meta": inst.meta,
    }
    (out / "meta.json").write_text(json.dumps(meta, indent=2))
    _write_rows(out / "E.csv", ["i", "k", "E"], ((*key, v) for key, v in inst.E.items()))
    _write_rows(out / "N.csv", ["j", "k", "N"], ((*key, v) for key, v in inst.N.items()))
    c_header = ["i", "j", "k", "C"] if inst.ship_cost_by_part else ["i", "j", "C"]
    _write_rows(out / "C.csv", c_header, ((*key, v) for key, v in inst.C.items()))
    _write_rows(out / "B.csv", ["m", "j", "k", "B"], ((*key, v) for key, v in inst.B.items()))
    return out


def read_instance_csv(in_dir) -> Instance:
    """Read an instance written by write_instance_csv. Ids come back as strings."""
    d = Path(in_dir)
    meta = json.loads((d / "meta.json").read_text())
    # Zeros dropped before validation (negatives survive and are rejected there).
    E = _drop_zeros({(r[0], r[1]): _num(r[2]) for r in _read_rows(d / "E.csv")})
    N = _drop_zeros({(r[0], r[1]): _num(r[2]) for r in _read_rows(d / "N.csv")})
    C = {tuple(r[:-1]): _num(r[-1]) for r in _read_rows(d / "C.csv")}
    B = {(r[0], r[1], r[2]): _num(r[3]) for r in _read_rows(d / "B.csv")}
    inst = Instance(
        sites=meta["sites"],
        parts=meta["parts"],
        suppliers=meta["suppliers"],
        E=E,
        N=N,
        C=C,
        B=B,
        UD=meta["UD"],
        FC1=meta["FC1"],
        FC2=meta["FC2"],
        name=meta.get("name", ""),
        meta=meta.get("meta", {}),
    )
    validate_instance(inst)
    return inst


def _num(s: str):
    v = float(s)
    return int(v) if v.is_integer() else v


def _write_rows(path: Path, header, rows) -> None:
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def _read_rows(path: Path):
    with open(path, newline="") as fh:
        r = csv.reader(fh)
        next(r)
        yield from r
